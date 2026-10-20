"""Unit and regression tests for deterministic-dual-reporter."""

import json
from pathlib import Path
import docx
from dual_reporter.models import TestRunManifest, DefectSeverity, TestStatus
from dual_reporter.html_builder import render_html_report
from dual_reporter.docx_builder import build_docx_report


def test_manifest_parsing():
    manifest_path = Path("examples/sample_manifest.json")
    with open(manifest_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    manifest = TestRunManifest.from_dict(data)

    assert manifest.run_id == "RUN-2026-E2E-09"
    assert manifest.total_cases == 10
    assert manifest.passed_cases == 8
    assert manifest.failed_cases == 2
    assert manifest.pass_rate == 80.0
    assert len(manifest.defects) == 2
    assert len(manifest.evidences) == 6


def test_evidence_sha256_generation():
    manifest_path = Path("examples/sample_manifest.json")
    with open(manifest_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    manifest = TestRunManifest.from_dict(data)

    for ev in manifest.evidences:
        assert len(ev.sha256_hash) == 64
        # Verify hex string
        int(ev.sha256_hash, 16)


def test_html_render():
    manifest_path = Path("examples/sample_manifest.json")
    with open(manifest_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    manifest = TestRunManifest.from_dict(data)
    html = render_html_report(manifest)

    assert "<!DOCTYPE html>" in html
    assert "RUN-2026-E2E-09" in html
    assert "Deadlock condition encountered" in html
    assert "efb3f7354fb3607f23f0ccedb992a82c5e5270d02a13bdb20dd91898d19e3204" in html


def test_docx_generation(tmp_path):
    manifest_path = Path("examples/sample_manifest.json")
    with open(manifest_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    manifest = TestRunManifest.from_dict(data)
    out_docx = tmp_path / "test_report.docx"
    build_docx_report(manifest, str(out_docx))

    assert out_docx.exists()
    doc = docx.Document(str(out_docx))
    
    # Assert headings and tables exist
    paragraphs_text = " ".join(p.text for p in doc.paragraphs)
    assert "Enterprise Core Gateway & Settlement Engine" in paragraphs_text
    assert "Executive Summary & Quality Scorecard" in paragraphs_text
    assert len(doc.tables) >= 4
